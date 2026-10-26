# -*- coding: utf-8 -*-
"""标书审查引擎：文档解析 + 规则审查
规则库源自 2026-08 摄像装置标书实战审查（17项问题归纳）。
@author WorkBuddy-GLM, Collus Wang
"""
import re
from pathlib import Path

# ---------------- 文档解析 ----------------

def parse_file(path):
    p = str(path)
    lower = p.lower()
    if lower.endswith(".pdf"):
        return parse_pdf(p)
    if lower.endswith(".docx"):
        return parse_docx(p)
    if lower.endswith(".doc"):
        return parse_doc(p)
    if lower.endswith((".txt", ".md")):
        text = Path(p).read_text("utf-8", errors="ignore")
        return _wrap([text])
    raise ValueError("暂不支持该文件类型（支持 PDF/DOCX/DOC/TXT）")


def parse_doc(path):
    """老版 .doc：调用本机 Microsoft Word（COM）提取文本。
    失败时抛出带明确指引的错误（Word 未安装 / 文件损坏 / 缺 pywin32），由上层写入摘要。"""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        raise ValueError("本机未安装 pywin32，无法解析 .doc。请将 .doc 另存为 docx 或 pdf 后重新上传。")
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        try:
            word.DisplayAlerts = 0
        except Exception:
            pass
        doc = word.Documents.Open(str(Path(path).resolve()), ReadOnly=True)
        try:
            text = doc.Content.Text
        finally:
            try:
                doc.Close(False)
            except Exception:
                pass
    except Exception as e:
        raise ValueError(
            f".doc 解析失败（请确认本机已安装 Microsoft Word，或将文件另存为 docx/pdf 后重新上传）：{e}")
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()
    # Word 文本以 \r 分段；去空行后按约3000字切伪页便于证据定位
    lines = [ln.strip() for ln in text.split("\r") if ln.strip()]
    pages, chunk = [], ""
    for ln in lines:
        chunk += ln + "\n"
        if len(chunk) >= 3000:
            pages.append(chunk)
            chunk = ""
    if chunk.strip():
        pages.append(chunk)
    if not pages:
        pages = [""]
    return _wrap(pages)


def parse_pdf(path):
    import fitz
    doc = fitz.open(path)
    pages = [page.get_text() for page in doc]
    n = len(doc)
    doc.close()
    return _wrap(pages)


def parse_docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    text = "\n".join(parts)
    # DOCX 无页概念，按约3000字切伪页便于证据定位
    pages, chunk = [], ""
    for line in text.split("\n"):
        chunk += line + "\n"
        if len(chunk) >= 3000:
            pages.append(chunk)
            chunk = ""
    if chunk.strip():
        pages.append(chunk)
    if not pages:
        pages = [""]
    return _wrap(pages)


def _wrap(pages):
    char_count = sum(len(p) for p in pages)
    return {
        "n_pages": len(pages),
        "pages": pages,
        "has_text_layer": char_count >= 100,
        "char_count": char_count,
    }

# ---------------- 工具 ----------------

def _line_of(text, pos):
    s = text.rfind("\n", 0, pos) + 1
    e = text.find("\n", pos)
    if e == -1:
        e = len(text)
    return text[s:e].strip()


def F(rule_id, category, severity, title, description, page, evidence, suggestion):
    return {
        "rule_id": rule_id, "category": category, "severity": severity,
        "title": title, "description": description,
        "page": page, "evidence": evidence, "suggestion": suggestion,
    }

# ---------------- 规则 ----------------

DUP_STOP = set(
    "天天 年年 人人 家家 时时 处处 常常 往往 渐渐 慢慢 轻轻 悄悄 刚刚 恰恰 "
    "白白 匆匆 谢谢 偏偏 层层 步步 细细 多多 等等 久久 遥遥 徐徐 缓缓 "
    "滔滔 源源 井井 津津 历历 依依 栩栩 亭亭 潺潺 皑皑 茫茫".split()
)

REQUIRED_SECTIONS_DEFAULT = [
    "报价一览表", "分项报价表", "商务条款偏离表", "技术规格偏离表",
    "产品生产进度计划", "售后服务", "同类项目业绩", "资格证明文件",
    "营业执照", "法定代表人身份证明", "法定代表人授权委托书",
    "质量管理体系", "技术方案",
]


def extract_required_from_text(text):
    """从招标要求文本中启发式提取"应提交资料"类名词清单（用于动态 R201）。
    仅在含"提交/提供/包括/包含/具备"等提示语的句子内提取，避免全文噪音。
    限制总长 ≤10 字、排除含连接词的句子型片段，保留短名词型资料名。
    提取不到或未提供文本时返回 []，此时 R201 不核对（不会把通用规则强加给不适用的标书）。"""
    if not text or not text.strip():
        return []
    segs = re.split(r"[。；;\n]", text)
    trig = re.compile(r"(应提交|需提交|应提供|需提供|必须提供|要求提供|提供以下|包括|包含|应具备|需具备|须提供|须提交|提交)")
    # 短名词型：2-8 字前缀 + 后缀，总长 ≤10
    doc_pat = re.compile(r"([\u4e00-\u9fa5]{2,8}(?:表|书|证明|证书|承诺|计划|方案|说明|清单|报告|执照|委托书|授权书))")
    NOISE = {"投标文件", "招标文件", "响应文件", "文件", "合同书", "中标通知书",
             "询比价文件", "报价文件", "采购方案", "方案要求", "询价文件"}
    # 句子型排除：含连接词/助词的长串
    BAD = re.compile(r"[的或和等如为以所及与]")
    out, seen = [], set()
    for seg in segs:
        if not trig.search(seg):
            continue
        for m in doc_pat.finditer(seg):
            w = m.group(1)
            if w in NOISE or w in seen:
                continue
            if len(w) > 10:
                continue
            if BAD.search(w):
                continue
            seen.add(w)
            out.append(w)
    return out


def rule_placeholders(bid):
    """R101 占位符未填写（XXXXXX、待填、下划线填空）"""
    pat = re.compile(r"[XxＸ×]{3,}|_{5,}|（待填|【待填|（待补充|【待补充|（空）")
    hits = []
    for i, p in enumerate(bid["pages"]):
        for m in pat.finditer(p):
            line = _line_of(p, m.start())
            if line:
                hits.append((i + 1, line[:80]))
    if not hits:
        return []
    ev = "\n".join(f"第{pg}页：{ln}" for pg, ln in hits[:10])
    if len(hits) > 10:
        ev += f"\n……共 {len(hits)} 处（仅列前10处）"
    return [F("R101", "格式", "high", f"发现 {len(hits)} 处占位符未填写",
              "标书中存在 XXXXXX/待填/空格填空等占位内容，若随标书提交将直接构成废标风险。",
              hits[0][0], ev, "逐处填写实际内容；提交前全局搜索 X、待填、下划线串做最终确认。")]


def rule_bookmark_error(bid):
    """R102 目录书签/引用错误"""
    pat = re.compile(r"错误!未定义书签|错误！未定义书签|错误!未找到引用源|错误！未找到引用源|Error! Bookmark")
    hits = []
    for i, p in enumerate(bid["pages"]):
        for m in pat.finditer(p):
            hits.append((i + 1, _line_of(p, m.start())[:80]))
    if not hits:
        return []
    ev = "\n".join(f"第{pg}页：{ln}" for pg, ln in hits[:10])
    return [F("R102", "格式", "high", f"目录存在 {len(hits)} 处书签错误",
              "Word 自动目录引用失效，导出 PDF 前未刷新域，评标专家会视为装订/编制质量问题。",
              hits[0][0], ev, "在 Word 中 Ctrl+A → F9 更新全部域后重新导出 PDF。")]


def rule_deviation_unmarked(bid):
    """R103 偏离表存在但无任何偏离标注"""
    full = "\n".join(bid["pages"])
    if "偏离表" not in full:
        return []
    n = len(re.findall(r"无偏离|正偏离|负偏离", full))
    if n == 0:
        return [F("R103", "技术", "high", "偏离表存在但未发现任何偏离标注",
                  "文档含偏离表，但全文没有出现“无偏离/正偏离/负偏离”标注，偏离列疑似漏填。",
                  0, "偏离表所在章节", "逐行填写响应/偏离结论；实际优于招标要求的应标“正偏离”以争取技术分。")]
    return []


def rule_price_blank(bid):
    """R104 报价表疑似未填写"""
    out = []
    # 总价行无数值
    for i, p in enumerate(bid["pages"]):
        for line in p.split("\n"):
            l = line.strip()
            if re.search(r"(竞选总价|投标总价|总报价|投标金额)", l) and not re.search(r"\d", l) and len(l) < 40:
                out.append((i + 1, l[:80]))
    if out:
        ev = "\n".join(f"第{pg}页：{ln}" for pg, ln in out[:6])
        out_issue = [F("R104", "完整性", "high", f"发现 {len(out)} 处总价栏疑似未填写",
                       "总价行存在但行内未检测到数字，报价表可能漏填。",
                       out[0][0], ev, "填写报价金额（含大小写），并与分项报价表合计核对一致。")]
        return out_issue
    # 报价一览表章节无金额
    for i, p in enumerate(bid["pages"]):
        if "报价一览表" in p:
            window = "\n".join(bid["pages"][i:i + 2])
            if not re.search(r"\d[\d,，.]*\s*元", window):
                return [F("R104", "完整性", "high", "报价一览表章节未检测到金额",
                          "找到“报价一览表”章节，但其后两页内无任何“数字+元”格式的金额。",
                          i + 1, "报价一览表所在页", "填写完整报价并核对大小写金额一致。")]
            break
    return []


def rule_schedule_thin(bid):
    """R105 生产进度计划疑似空白/过于简略"""
    for i, p in enumerate(bid["pages"]):
        if "生产进度计划" in p:
            window = "\n".join(bid["pages"][i:i + 3])
            dates = re.findall(r"\d{1,4}\s*[年月日/\-]\s*\d{1,2}", window)
            if len(dates) < 3:
                return [F("R105", "技术", "medium", "产品生产进度计划疑似空白或过于简略",
                          "找到“生产进度计划”章节，但其后内容中时间节点过少（少于3处），可能只有标题没有计划表。",
                          i + 1, "生产进度计划章节", "补充按周/阶段划分的进度表：设计→采购→生产→测试→交付，每个节点标注天数或日期。")]
            return []
    return []


def rule_duplicate_chars(bid):
    """R106 疑似重复字（如“包括包括”）"""
    pat = re.compile(r"([\u4e00-\u9fa5]{2})\1")
    hits, seen = [], set()
    for i, p in enumerate(bid["pages"]):
        for m in pat.finditer(p):
            word = m.group(0)
            base = m.group(1)
            if base in DUP_STOP or word in seen:
                continue
            seen.add(word)
            hits.append((i + 1, word, _line_of(p, m.start())[:60]))
    if not hits:
        return []
    ev = "\n".join(f"第{pg}页：“{w}”｜{ln}" for pg, w, ln in hits[:12])
    if len(hits) > 12:
        ev += f"\n……共 {len(hights) if False else len(hits)} 处（仅列前12处）"
    return [F("R106", "格式", "low", f"发现 {len(hits)} 处疑似重复字",
              "两字连续重复（如“包括包括”），多为复制粘贴或删改残留，属低级文字错误。",
              hits[0][0], ev, "逐处确认；叠词除外（如“天天”）均为笔误，需删除重复字。")]


def rule_signature_blank(bid):
    """R107 签字/盖章/日期栏疑似空白"""
    pat = re.compile(r"(法定代表人|授权代表|被授权人|投标人|委托代理人|供应商)[^\n]{0,12}(签字|盖章|印鉴)[:：]?\s*$|^\s*(签字|盖章)[:：]?\s*$|日期[:：]?\s*$", re.M)
    hits = []
    for i, p in enumerate(bid["pages"]):
        for m in pat.finditer(p):
            line = m.group(0).strip()
            if line:
                hits.append((i + 1, line[:60]))
    if not hits:
        return []
    ev = "\n".join(f"第{pg}页：{ln}" for pg, ln in hits[:8])
    return [F("R107", "格式", "medium", f"发现 {len(hits)} 处签字/盖章/日期栏疑似空白",
              "签字、盖章或日期行后无内容。注意：PDF文本层无法100%确认是否已签（可能为手写/图片签章），需人工复核。",
              hits[0][0], ev, "打印后逐页补签字盖章；授权委托书务必双签（法定代表人+被授权人）。")]


def rule_version_prefix(bid):
    """R108 版本号疑似缺 V 前缀"""
    pat = re.compile(r"版本号[:：]?\s*(?![VvＶ])[0-9][0-9A-Za-z.]{0,8}")
    hits = []
    for i, p in enumerate(bid["pages"]):
        for m in pat.finditer(p):
            hits.append((i + 1, _line_of(p, m.start())[:70]))
    if not hits:
        return []
    ev = "\n".join(f"第{pg}页：{ln}" for pg, ln in hits[:6])
    return [F("R108", "格式", "low", f"发现 {len(hits)} 处版本号疑似缺“V”前缀",
              "版本号写为“3.0”而非“V3.0”，与标书其他部分（或招标文件）的写法可能不一致。",
              hits[0][0], ev, "统一为“V3.0”格式，并与技术方案、偏离表中的版本号写法保持一致。")]


def rule_amount_consistency(bid):
    """R109 总价金额多处不一致"""
    amts = {}
    for i, p in enumerate(bid["pages"]):
        for line in p.split("\n"):
            if re.search(r"竞选总价|投标总价|总报价|投标金额", line):
                for m in re.finditer(r"([0-9][0-9,，.]{2,15})\s*元?", line):
                    val = m.group(1).replace(",", "").replace("，", "")
                    try:
                        v = float(val)
                    except ValueError:
                        continue
                    amts.setdefault(v, []).append(i + 1)
    if len(amts) >= 2:
        ev = "\n".join(f"{v:,.2f} 元 → 第{'、'.join(map(str, pgs[:4]))}页" for v, pgs in sorted(amts.items()))
        return [F("R109", "技术", "high", f"总价出现 {len(amts)} 个不同数值",
                  "响应文件、报价一览表、分项报价表等处的总价应完全一致，不一致会被认定为报价矛盾。",
                  0, ev, "统一各处总价；检查是否含税、是否含运费口径不一致导致。")]
    return []


def rule_required_sections(bid, required, source_label="招标要求"):
    """R201 必备章节/资料缺失（清单来自招标要求动态提取或用户配置）
    匹配容忍“响应/偏离表”等变体措辞：去斜杠/空格后匹配，或去“表”字匹配。"""
    full = "\n".join(bid["pages"])
    # 匹配容忍变体措辞：去斜杠/顿号/空格；再去“响应”二字（应对“商务条款响应/偏离表”类写法）
    variants = [
        full,
        re.sub(r"[/／、\s]", "", full),
        re.sub(r"[/／、\s]", "", full).replace("响应", ""),
    ]
    out = []
    for kw in required:
        if not kw:
            continue
        base = kw.rstrip("表")
        if any(kw in v or (base and base in v) for v in variants):
            continue
        out.append(F("R201", "完整性", "high", f"缺少必备内容：{kw}",
                     f"招标要求（{source_label}）明确列出的资料，标书中未检索到。必备资料缺失是常见废标原因。",
                     0, f"全文搜索“{kw}”无结果（已容忍“响应/偏离表”等变体措辞）",
                     f"补充“{kw}”相关章节或证明文件；若已有但换了名称，请确认标题与招标要求措辞一致。"))
    return out

# ---------------- 主入口 ----------------

def run_rules(bid, config, tender=None, required_from_rules=None, tender_error=None):
    issues = []
    summary = {
        "标书页数": bid["n_pages"],
        "标书字符数": bid["char_count"],
        "标书文本层": "有" if bid["has_text_layer"] else "无（扫描件）",
    }
    if not bid["has_text_layer"]:
        issues.append(F("R000", "完整性", "high", "标书疑似扫描件，无法自动审查",
                        "未检测到文本层（纯图片PDF），规则引擎与LLM均无法读取内容。",
                        0, "", "上传文字版PDF/DOCX；或先做OCR转换后再上传。"))
        return issues, summary

    if tender:
        if tender["has_text_layer"]:
            summary["招标文件"] = f"已解析（{tender['n_pages']}页），将自动逐项核对"
        else:
            summary["招标文件"] = "已上传但为扫描件，自动核对受限"
            issues.append(F("R202", "完整性", "low", "招标文件为扫描件，自动核对受限",
                            "上传的招标文件无文本层，系统无法自动提取其要求做逐项核对。可在“补充检查项”中粘贴关键要求（必备资料清单/技术参数），即可获得针对性审查。",
                            0, "", "在“补充检查项”中粘贴招标文件的关键要求，或上传文本版招标文件。"))
    elif tender_error:
        summary["招标文件"] = f"解析失败：{tender_error}"
        issues.append(F("R202", "完整性", "medium", "招标文件解析失败，交叉核对已跳过",
                        f"上传的招标文件无法解析：{tender_error}。请转为 PDF/DOCX/TXT 后重新上传，或在“补充检查项”中粘贴关键要求。",
                        0, "", "将 .doc 另存为 docx/pdf 后重新上传，或在“补充检查项”中粘贴招标文件的关键要求。"))
    else:
        issues.append(F("R202", "完整性", "low", "未上传招标文件，交叉核对已跳过",
                        "仅基于标书自身与内置通用规则审查。上传招标文件（文本版）后可自动逐项核对；也可在“补充检查项”中填写重点关注项。",
                        0, "", "上传招标文件（PDF/DOCX/DOC/TXT，需含文本层），或在“补充检查项”中填写关键要求。"))

    # R201 必备资料核对：优先用招标要求动态提取的清单；其次用户全局配置；都没有则跳过（不误报）
    required = required_from_rules if required_from_rules else config.get("required_sections", [])
    if required:
        issues += rule_required_sections(bid, required, source_label="招标要求")
        summary["必备资料核对"] = f"按招标要求核对 {len(required)} 项：{'、'.join(required)}"
    else:
        summary["必备资料核对"] = "未提供招标要求文本，跳过必备资料核对（避免将不适用规则强加给本标书）"

    # 通用规则：与具体招标要求无关的基础质量检查，任何标书都适用
    issues += rule_placeholders(bid)
    issues += rule_bookmark_error(bid)
    issues += rule_deviation_unmarked(bid)
    issues += rule_price_blank(bid)
    issues += rule_schedule_thin(bid)
    issues += rule_amount_consistency(bid)
    issues += rule_signature_blank(bid)
    issues += rule_duplicate_chars(bid)
    issues += rule_version_prefix(bid)

    return issues, summary
